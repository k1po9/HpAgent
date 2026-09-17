"""Bounded PPTX reads using python-pptx."""

from __future__ import annotations

from typing import Any

from file_domain.models import FileResource


class PptxAdapter:
    @staticmethod
    def _presentation(resource: FileResource):
        from pptx import Presentation

        return Presentation(resource.local_path)

    def inspect(self, resource: FileResource) -> dict[str, Any]:
        presentation = self._presentation(resource)
        return {"slides": len(presentation.slides)}

    def read_slide(
        self, resource: FileResource, slide: int, max_chars: int
    ) -> dict[str, Any]:
        presentation = self._presentation(resource)
        if slide < 1 or slide > len(presentation.slides):
            raise ValueError("invalid presentation slide")
        texts: list[str] = []
        used = 0
        truncated = False
        for shape in presentation.slides[slide - 1].shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            value = str(shape.text or "")
            remaining = max_chars - used
            if remaining <= 0:
                truncated = True
                break
            texts.append(value[:remaining])
            used += len(texts[-1])
            if len(texts[-1]) < len(value):
                truncated = True
                break
        return {"slide": slide, "texts": texts, "truncated": truncated}
