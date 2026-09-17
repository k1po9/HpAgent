"""Declarative Office output adapters; never execute user-supplied code."""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from file_domain.models import FileResource


def _atomic_target(outputs_root: Path, logical_name: str, suffix: str) -> tuple[Path, Path]:
    if not logical_name.casefold().endswith(suffix):
        raise ValueError(f"output name must end with {suffix}")
    target = (outputs_root / logical_name).absolute()
    if not target.is_relative_to(outputs_root) or target.parent != outputs_root:
        raise ValueError("output escapes Run output scope")
    temporary = outputs_root / f".{logical_name}.part"
    if target.exists() or temporary.exists():
        raise FileExistsError("output path already exists")
    return temporary, target


def _commit(temporary: Path, target: Path) -> Path:
    with temporary.open("rb") as stream:
        os.fsync(stream.fileno())
    os.replace(temporary, target)
    return target


class DocxWriter:
    def create(self, outputs_root: Path, logical_name: str, blocks: list[dict[str, Any]]) -> Path:
        from docx import Document

        temporary, target = _atomic_target(outputs_root, logical_name, ".docx")
        try:
            document = Document()
            for block in blocks:
                kind = block.get("kind", "paragraph")
                text = str(block.get("text", ""))
                if kind == "heading":
                    document.add_heading(text, level=int(block.get("level", 1)))
                elif kind == "paragraph":
                    document.add_paragraph(text)
                else:
                    raise ValueError("unsupported DOCX block kind")
            document.save(temporary)
            return _commit(temporary, target)
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise

    def replace_text(
        self, resource: FileResource, outputs_root: Path, logical_name: str,
        find: str, replace: str, max_replacements: int,
    ) -> tuple[Path, int]:
        from docx import Document

        if not find:
            raise ValueError("find text must not be empty")
        temporary, target = _atomic_target(outputs_root, logical_name, ".docx")
        try:
            document = Document(resource.local_path)
            count = 0
            for paragraph in document.paragraphs:
                remaining = max_replacements - count
                if remaining <= 0:
                    break
                occurrences = min(paragraph.text.count(find), remaining)
                if occurrences:
                    paragraph.text = paragraph.text.replace(find, replace, occurrences)
                    count += occurrences
            if count == 0:
                raise ValueError("DOCX text was not found")
            document.save(temporary)
            return _commit(temporary, target), count
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise

    def append_section(
        self, resource: FileResource, outputs_root: Path, logical_name: str,
        heading: str, paragraphs: list[str],
    ) -> Path:
        from docx import Document

        temporary, target = _atomic_target(outputs_root, logical_name, ".docx")
        try:
            document = Document(resource.local_path)
            if heading:
                document.add_heading(heading, level=1)
            for text in paragraphs:
                document.add_paragraph(text)
            document.save(temporary)
            return _commit(temporary, target)
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise


class XlsxWriter:
    def create(
        self, outputs_root: Path, logical_name: str, sheets: list[dict[str, Any]]
    ) -> Path:
        from openpyxl import Workbook

        temporary, target = _atomic_target(outputs_root, logical_name, ".xlsx")
        names = [str(sheet["name"]) for sheet in sheets]
        if len({name.casefold() for name in names}) != len(names):
            raise ValueError("workbook sheet names must be unique")
        try:
            workbook = Workbook()
            workbook.remove(workbook.active)
            for sheet in sheets:
                worksheet = workbook.create_sheet(str(sheet["name"]))
                for row in sheet.get("rows", []):
                    worksheet.append(list(row))
            workbook.save(temporary)
            return _commit(temporary, target)
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise

    def write_range(
        self, resource: FileResource, outputs_root: Path, logical_name: str,
        sheet: str, start_row: int, start_column: int, values: list[list[Any]],
    ) -> Path:
        from openpyxl import load_workbook

        temporary, target = _atomic_target(outputs_root, logical_name, ".xlsx")
        try:
            workbook = load_workbook(resource.local_path)
            if sheet not in workbook.sheetnames:
                raise ValueError("unknown workbook sheet")
            worksheet = workbook[sheet]
            for row_offset, row in enumerate(values):
                for column_offset, value in enumerate(row):
                    worksheet.cell(start_row + row_offset, start_column + column_offset, value)
            workbook.save(temporary)
            return _commit(temporary, target)
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise


class PptxWriter:
    def create(
        self, outputs_root: Path, logical_name: str, slides: list[dict[str, str]]
    ) -> Path:
        from pptx import Presentation

        temporary, target = _atomic_target(outputs_root, logical_name, ".pptx")
        try:
            presentation = Presentation()
            for slide_data in slides:
                slide = presentation.slides.add_slide(presentation.slide_layouts[1])
                slide.shapes.title.text = slide_data.get("title", "")
                slide.placeholders[1].text = slide_data.get("body", "")
            presentation.save(temporary)
            return _commit(temporary, target)
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise

    def replace_slide(
        self, resource: FileResource, outputs_root: Path, logical_name: str,
        slide_number: int, title: str, body: str,
    ) -> Path:
        from pptx import Presentation

        temporary, target = _atomic_target(outputs_root, logical_name, ".pptx")
        try:
            presentation = Presentation(resource.local_path)
            if slide_number < 1 or slide_number > len(presentation.slides):
                raise ValueError("invalid presentation slide")
            slide = presentation.slides[slide_number - 1]
            title_shape = slide.shapes.title
            if title_shape is not None:
                title_shape.text = title
            body_shape = next(
                (
                    shape for shape in slide.placeholders
                    if getattr(shape, "has_text_frame", False)
                    and (
                        title_shape is None
                        or shape._element is not title_shape._element
                    )
                ),
                None,
            )
            if body_shape is None:
                raise ValueError("slide has no editable body placeholder")
            body_shape.text = body
            presentation.save(temporary)
            return _commit(temporary, target)
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise
