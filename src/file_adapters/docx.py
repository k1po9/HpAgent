"""Bounded DOCX reads using python-docx."""

from __future__ import annotations

from typing import Any

from file_domain.models import FileResource


class DocxAdapter:
    @staticmethod
    def _document(resource: FileResource):
        from docx import Document

        return Document(resource.local_path)

    def inspect(self, resource: FileResource) -> dict[str, Any]:
        document = self._document(resource)
        sections = [
            {"index": index, "heading": paragraph.text.strip()}
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph.style and paragraph.style.name.startswith("Heading")
        ]
        return {
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "sections": sections,
        }

    def read_paragraphs(
        self, resource: FileResource, start: int, limit: int, max_chars: int
    ) -> dict[str, Any]:
        document = self._document(resource)
        if start < 0 or limit < 1:
            raise ValueError("invalid paragraph range")
        selected = document.paragraphs[start : start + limit]
        result: list[dict[str, Any]] = []
        used = 0
        truncated = start + limit < len(document.paragraphs)
        for index, paragraph in enumerate(selected, start=start):
            remaining = max_chars - used
            if remaining <= 0:
                truncated = True
                break
            text = paragraph.text[:remaining]
            result.append({"index": index, "style": paragraph.style.name, "text": text})
            used += len(text)
            if len(text) < len(paragraph.text):
                truncated = True
                break
        return {"paragraphs": result, "truncated": truncated}

    def extract_tables(
        self, resource: FileResource, table_index: int, max_rows: int
    ) -> dict[str, Any]:
        document = self._document(resource)
        if table_index < 0 or table_index >= len(document.tables):
            raise ValueError("invalid DOCX table index")
        rows = [
            [cell.text for cell in row.cells]
            for row in document.tables[table_index].rows[:max_rows]
        ]
        return {
            "table_index": table_index,
            "rows": rows,
            "truncated": len(document.tables[table_index].rows) > max_rows,
        }
