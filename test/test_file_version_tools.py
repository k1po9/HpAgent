from __future__ import annotations

import json
from uuid import uuid4

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from file_runtime import PublishedOutput
from sandbox.tools.local.file_write import create_file_write_tools
from workspace.file_scope import RunFileInput, RunFileScope


class _Publisher:
    def replay(self, scope, operation_id, logical_name, parent_file_id=None):
        return None

    def publish(
        self, scope, operation_id, logical_name, content_type, *, parent_file_id=None
    ):
        path = scope.outputs_root / logical_name
        return PublishedOutput(
            uuid4(), logical_name, content_type, path.stat().st_size, "b" * 64,
            parent_file_id, 2,
        )


@pytest.mark.asyncio
async def test_replace_docx_text_creates_child_version(tmp_path):
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    document = Document()
    document.add_paragraph("old value")
    document.save(inputs / "source.docx")
    for name in ("outputs", "scratch"):
        (tmp_path / name).mkdir()
    source_id = uuid4()
    scope = RunFileScope(
        uuid4(), inputs, tmp_path / "scratch", tmp_path / "outputs",
        (RunFileInput(source_id, "source.docx", (inputs / "source.docx").stat().st_size,
                      "binary", "application/docx"),),
    )
    tools = {tool.name: tool for tool in create_file_write_tools(lambda: scope, _Publisher())}
    result = json.loads(await tools["replace_docx_text"].ainvoke({
        "file": "source.docx", "output_name": "source-v2.docx",
        "find": "old", "replace": "new", "operation_id": "patch:docx",
    }))
    assert result["parent_file_id"] == str(source_id)
    assert result["version"] == 2
    assert result["replacements"] == 1
    assert Document(scope.outputs_root / "source-v2.docx").paragraphs[0].text == "new value"
    appended = json.loads(await tools["append_docx_section"].ainvoke({
        "file": "source.docx", "output_name": "source-append.docx",
        "heading": "Added", "paragraphs": ["bounded paragraph"],
        "operation_id": "patch:append",
    }))
    saved = Document(scope.outputs_root / "source-append.docx")
    assert saved.paragraphs[-1].text == "bounded paragraph"
    assert appended["parent_file_id"] == str(source_id)


@pytest.mark.asyncio
async def test_write_sheet_range_creates_child_without_formula_engine(tmp_path):
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.save(inputs / "source.xlsx")
    for name in ("outputs", "scratch"):
        (tmp_path / name).mkdir()
    source_id = uuid4()
    scope = RunFileScope(
        uuid4(), inputs, tmp_path / "scratch", tmp_path / "outputs",
        (RunFileInput(source_id, "source.xlsx", (inputs / "source.xlsx").stat().st_size,
                      "binary", "application/xlsx"),),
    )
    tools = {tool.name: tool for tool in create_file_write_tools(lambda: scope, _Publisher())}
    result = json.loads(await tools["write_sheet_range"].ainvoke({
        "file": "source.xlsx", "output_name": "source-v2.xlsx", "sheet": "Data",
        "start_row": 2, "start_column": 2, "values": [["x", "=1+1"]],
        "operation_id": "patch:xlsx",
    }))
    saved = load_workbook(scope.outputs_root / "source-v2.xlsx", data_only=False)
    assert saved["Data"]["B2"].value == "x"
    assert saved["Data"]["C2"].value == "=1+1"
    assert result["parent_file_id"] == str(source_id)


@pytest.mark.asyncio
async def test_replace_slide_creates_child_version(tmp_path):
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "old"
    slide.placeholders[1].text = "body"
    presentation.save(inputs / "source.pptx")
    for name in ("outputs", "scratch"):
        (tmp_path / name).mkdir()
    source_id = uuid4()
    scope = RunFileScope(
        uuid4(), inputs, tmp_path / "scratch", tmp_path / "outputs",
        (RunFileInput(source_id, "source.pptx", (inputs / "source.pptx").stat().st_size,
                      "binary", "application/pptx"),),
    )
    tools = {tool.name: tool for tool in create_file_write_tools(lambda: scope, _Publisher())}
    result = json.loads(await tools["replace_slide"].ainvoke({
        "file": "source.pptx", "output_name": "source-v2.pptx", "slide": 1,
        "title": "new", "body": "updated", "operation_id": "patch:pptx",
    }))
    saved = Presentation(scope.outputs_root / "source-v2.pptx")
    assert saved.slides[0].shapes.title.text == "new"
    assert result["version"] == 2
