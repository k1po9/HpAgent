from __future__ import annotations

import io
import json
from uuid import UUID, uuid4

import pytest
from docx import Document

from conversation_domain.commands import CommandService
from file_runtime import FileResourceResolver, OutputPublisher
from research_domain.services import ResearchTaskCommandService
from sandbox.git_repo import GitRepoManager
from sandbox.tools.local.file_write import create_file_write_tools
from storage.tenant_file_store import TenantFileStore
from web_domain.file_cleanup import FileCleanupService
from web_domain.file_services import FileService
from workspace.catalog import WorkspaceCatalog, WorkspaceVersionConflict
from workspace.file_scope import RunFileWorkspace
from workspace.resources import ResourcePolicy

pytestmark = [pytest.mark.postgres, pytest.mark.asyncio]


async def test_real_docx_upload_edit_revision_and_conflict_save(
    db, account_id, database_url, worker_database_url, tmp_path, monkeypatch,
):
    def no_git(*args, **kwargs):
        raise AssertionError("ordinary Workspace files invoked GitRepoManager")

    monkeypatch.setattr(GitRepoManager, "repo_path", no_git)
    monkeypatch.setattr(GitRepoManager, "ensure_session_workspace", no_git)
    store = TenantFileStore(tmp_path / "store", max_bytes=1024 * 1024)
    files = FileService(database_url, store, max_bytes=1024 * 1024)
    catalog = WorkspaceCatalog(database_url)
    commands = CommandService(database_url)
    policy = ResourcePolicy(database_url)
    tree = catalog.initialize(account_id)
    parent = UUID(next(n["node_id"] for n in tree["nodes"] if n["name"] == "资料"))
    document = Document()
    document.add_paragraph("revision one")
    initial = io.BytesIO()
    document.save(initial)
    body = initial.getvalue()
    created = files.create_workspace_upload(account_id, str(uuid4()), "draft.docx",
                                            len(body), "application/vnd.openxmlformats-officedocument."
                                            "wordprocessingml.document", None)
    first_file = UUID(created["file"]["file_id"])

    async def chunks():
        yield body

    await files.upload_content(account_id, first_file, chunks())
    entry = catalog.save_file(account_id, parent, first_file, "draft.docx", "save:docx")
    catalog.upgrade_file(account_id, entry)
    baseline = catalog.versions(account_id, entry)["current"]

    def authorized_run() -> UUID:
        conversation = UUID(commands.create_conversation(
            account_id, str(uuid4()))["conversation_id"])
        policy.grant(account_id, "conversation", conversation, entry,
                     ["list_metadata", "read_content", "update_content"], False)
        run_id = UUID(commands.send_message(account_id, conversation,
                                            str(uuid4()), "edit docx")["run_id"])
        assert policy.select(account_id, run_id, entry)["file_id"] == str(first_file)
        return run_id

    old_run = authorized_run()
    edit_run = authorized_run()
    tasks = ResearchTaskCommandService(database_url)
    task_id = UUID(tasks.create_task(account_id, str(uuid4()), "Reader", "Read DOCX")
                   .body["task_id"])
    policy.grant(account_id, "task", task_id, entry,
                 ["list_metadata", "read_content"], False)
    task_run = UUID(tasks.trigger_task(account_id, task_id, str(uuid4())).body["run_id"])
    assert policy.select(account_id, task_run, entry)["file_id"] == str(first_file)
    publisher = OutputPublisher(worker_database_url, store)
    workspace = RunFileWorkspace(worker_database_url, store, tmp_path / "execution")
    with workspace.prepare(account_id, edit_run, include_selected=True) as scope:
        tools = {tool.name: tool for tool in create_file_write_tools(lambda: scope, publisher)}
        changed = json.loads(await tools["replace_docx_text"].ainvoke({
            "file": "draft.docx", "output_name": "draft-v2.docx",
            "find": "one", "replace": "two", "operation_id": "edit:docx:v2",
        }))
        conflict = json.loads(await tools["replace_docx_text"].ainvoke({
            "file": "draft.docx", "output_name": "draft-alternate.docx",
            "find": "one", "replace": "alternate", "operation_id": "edit:docx:alternate",
        }))
    revised_file = UUID(changed["file_id"])
    conflict_file = UUID(conflict["file_id"])
    result = catalog.update_file(account_id, entry, edit_run, revised_file, 1,
                                 baseline["sha256"], "version:docx:v2")
    assert result["revision"] == 2
    assert Document(files.download(account_id, revised_file)[1]).paragraphs[0].text == "revision two"
    with workspace.prepare(account_id, old_run, include_selected=True) as scope:
        old = FileResourceResolver(scope).resolve(str(first_file))
        assert Document(old.local_path).paragraphs[0].text == "revision one"
    with workspace.prepare(account_id, task_run, include_selected=True) as scope:
        assert Document(FileResourceResolver(scope).resolve(str(first_file)).local_path) \
            .paragraphs[0].text == "revision one"
    assert db.execute("SELECT fixed_revision FROM run_resource_candidates "
                      "WHERE run_id=%s AND node_id=%s", (old_run, entry)).fetchone()[0] == 1
    with pytest.raises(WorkspaceVersionConflict):
        catalog.update_file(account_id, entry, edit_run, conflict_file, 1,
                            baseline["sha256"], "version:docx:conflict")
    alternate = catalog.save_file(account_id, parent, conflict_file,
                                  "draft-alternate.docx", "save:docx:conflict")
    assert catalog.versions(account_id, alternate)["current"]["file_id"] == str(conflict_file)
    assert Document(files.download(account_id, conflict_file)[1]).paragraphs[0].text == (
        "revision alternate"
    )
    # A separate direct upload has its long-lived reference released and its
    # bytes collected without creating a Git workspace.
    disposable = files.create_workspace_upload(account_id, str(uuid4()), "gc.txt",
                                               2, "text/plain", None)
    disposable_id = UUID(disposable["file"]["file_id"])

    async def disposable_chunks():
        yield b"gc"

    await files.upload_content(account_id, disposable_id, disposable_chunks())
    disposable_entry = catalog.save_file(account_id, parent, disposable_id,
                                         "gc.txt", "save:gitless-gc")
    catalog.remove(account_id, disposable_entry)
    files.delete(account_id, disposable_id)
    assert FileCleanupService(worker_database_url, store).cleanup_once().deleted == 1
    assert db.execute("SELECT storage_key FROM stored_files WHERE file_id=%s",
                      (disposable_id,)).fetchone()[0] is None
    assert not list(tmp_path.rglob(".git"))
