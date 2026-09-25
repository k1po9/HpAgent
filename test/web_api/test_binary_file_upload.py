from __future__ import annotations

import hashlib
import json
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter

from sandbox.tools.local.file_analysis import create_file_analysis_tools
from sandbox.tools.local.file_read import create_file_read_tools
from storage.tenant_file_store import TenantFileReader
from workspace.file_scope import RunFileWorkspace


def _headers(csrf: str, key: str | None = None) -> dict[str, str]:
    headers = {"Origin": "https://testserver", "X-CSRF-Token": csrf}
    if key:
        headers["Idempotency-Key"] = key
    return headers


def _login(client, username: str) -> str:
    response = client.post(
        "/auth/login",
        json={"username": username, "password": "correct-password", "return_to": "/"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    return str(client.get("/api/v1/me").json()["csrf_token"])


def _fixtures(root: Path) -> dict[str, tuple[str, bytes]]:
    document = Document()
    document.add_paragraph("binary docx")
    document.save(root / "sample.docx")
    workbook = Workbook()
    workbook.active.append(["binary", "xlsx"])
    workbook.save(root / "sample.xlsx")
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[5])
    presentation.save(root / "sample.pptx")
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with (root / "sample.pdf").open("wb") as stream:
        writer.write(stream)
    return {
        "sample.pdf": ("application/pdf", (root / "sample.pdf").read_bytes()),
        "sample.docx": (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            (root / "sample.docx").read_bytes(),
        ),
        "sample.xlsx": (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            (root / "sample.xlsx").read_bytes(),
        ),
        "sample.pptx": (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            (root / "sample.pptx").read_bytes(),
        ),
        "events.log": ("text/x-log", b"alpha event\nbeta event\nalpha done\n"),
        "笔记.md": ("text/markdown", "# 标题\n中文内容。\n".encode("utf-8")),
        "note.txt": ("text/plain", b"plain text\n"),
    }


@pytest.mark.asyncio
async def test_real_web_upload_binds_binary_and_text_files_into_run_scope(
    tmp_path, seed_identity, client_factory, worker_database_url,
):
    account_id = seed_identity("binary-user")
    store_root = tmp_path / "store"
    client = client_factory(
        file_upload_enabled=True, file_store_root=str(store_root)
    )
    csrf = _login(client, "binary-user")
    conversation_id = client.post(
        "/api/v1/conversations", json={"title": None},
        headers=_headers(csrf, str(uuid4())),
    ).json()["conversation"]["conversation_id"]
    uploaded: dict[str, UUID] = {}
    fixtures = _fixtures(tmp_path)

    for name, (media_type, content) in fixtures.items():
        created = client.post(
            f"/api/v1/conversations/{conversation_id}/uploads",
            json={
                "file_name": name,
                "size_bytes": len(content),
                "content_type": media_type,
                "sha256": hashlib.sha256(content).hexdigest(),
            },
            headers=_headers(csrf, str(uuid4())),
        )
        assert created.status_code == 201
        file_id = UUID(created.json()["file"]["file_id"])
        uploaded[name] = file_id
        completed = client.put(
            f"/api/v1/uploads/{file_id}/content",
            content=content,
            headers={
                **_headers(csrf),
                "Content-Type": "application/octet-stream",
            },
        )
        assert completed.status_code == 200
        expected_encoding = "utf-8" if name.endswith((".log", ".md", ".txt")) else "binary"
        assert completed.json()["file"]["encoding"] == expected_encoding
        assert completed.json()["file"]["content_type"] == media_type

    sent = client.post(
        f"/api/v1/conversations/{conversation_id}/messages",
        json={
            "content": "inspect attachments",
            "file_ids": [str(value) for value in uploaded.values()],
        },
        headers=_headers(csrf, str(uuid4())),
    )
    assert sent.status_code == 202
    run_id = UUID(sent.json()["run"]["run_id"])
    workspace = RunFileWorkspace(
        worker_database_url, TenantFileReader(store_root), tmp_path / "runs"
    )
    with workspace.prepare(account_id, run_id) as scope:
        assert {item.logical_name for item in scope.inputs} == set(fixtures)
        for item in scope.inputs:
            assert (scope.inputs_root / item.logical_name).read_bytes() == fixtures[item.logical_name][1]
        tools = {
            tool.name: tool for tool in create_file_analysis_tools(lambda: scope)
        }
        searched = json.loads(await tools["search_file"].ainvoke({
            "file": "events.log", "query": "alpha", "need_total": True,
        }))
        counted = json.loads(await tools["count_matches"].ainvoke({
            "file": "events.log", "query": "alpha",
        }))
        stats = json.loads(await tools["text_stats"].ainvoke({
            "file": "events.log", "keywords": ["alpha", "beta"],
        }))
        assert searched["total_matches"] == counted["count"] == 2
        assert stats["total_lines"] == 3
        read_tools = {tool.name: tool for tool in create_file_read_tools(lambda: scope)}
        markdown = json.loads(await read_tools["read_file"].ainvoke({"file": "笔记.md"}))
        assert "中文内容" in markdown["text"]
        assert markdown["metadata"]["adapter"] == "markitdown"


@pytest.mark.asyncio
async def test_markdown_alias_empty_mime_encoding_and_upload_guards(
    tmp_path, seed_identity, client_factory,
):
    seed_identity("markdown-owner")
    seed_identity("markdown-other")
    client = client_factory(file_upload_enabled=True, file_store_root=str(tmp_path / "store"))
    csrf = _login(client, "markdown-owner")
    conversation_id = client.post(
        "/api/v1/conversations", json={"title": None},
        headers=_headers(csrf, str(uuid4())),
    ).json()["conversation"]["conversation_id"]

    for media_type in ("text/x-markdown", ""):
        content = "# 中文\n".encode("utf-8")
        created = client.post(
            f"/api/v1/conversations/{conversation_id}/uploads",
            json={"file_name": "notes.md", "size_bytes": len(content), "content_type": media_type},
            headers=_headers(csrf, str(uuid4())),
        )
        assert created.status_code == 201
        assert created.json()["file"]["content_type"] == "text/markdown"
        file_id = created.json()["file"]["file_id"]
        uploaded = client.put(f"/api/v1/uploads/{file_id}/content", content=content,
                              headers={**_headers(csrf), "Content-Type": "application/octet-stream"})
        assert uploaded.status_code == 200
        assert uploaded.json()["file"]["encoding"] == "utf-8"

    invalid = b"\xff\xfe"
    created = client.post(
        f"/api/v1/conversations/{conversation_id}/uploads",
        json={"file_name": "broken.md", "size_bytes": len(invalid),
              "content_type": "text/markdown"},
        headers=_headers(csrf, str(uuid4())),
    )
    file_id = created.json()["file"]["file_id"]
    rejected = client.put(f"/api/v1/uploads/{file_id}/content", content=invalid,
                          headers={**_headers(csrf), "Content-Type": "application/octet-stream"})
    assert rejected.status_code == 422
    assert rejected.json()["error"]["code"] == "file_encoding_unsupported"
    assert client.get(f"/api/v1/files/{file_id}").json()["file"]["failure_code"] == "file_encoding_unsupported"

    unsupported = client.post(
        f"/api/v1/conversations/{conversation_id}/uploads",
        json={"file_name": "bad.exe", "size_bytes": 1, "content_type": "application/x-msdownload"},
        headers=_headers(csrf, str(uuid4())),
    )
    assert unsupported.status_code == 415
    oversized = client.post(
        f"/api/v1/conversations/{conversation_id}/uploads",
        json={"file_name": "large.md", "size_bytes": 128 * 1024 * 1024 + 1,
              "content_type": "text/markdown"},
        headers=_headers(csrf, str(uuid4())),
    )
    assert oversized.status_code == 413

    _login(client, "markdown-other")
    assert client.get(f"/api/v1/files/{file_id}").status_code == 404
