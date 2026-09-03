from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter

from file_runtime import FileAdapterRegistry
from sandbox.tools.local.file_read import create_file_read_tools
from workspace.file_scope import RunFileInput, RunFileScope


def _scope(tmp_path: Path) -> RunFileScope:
    inputs = tmp_path / "inputs"
    inputs.mkdir()

    document = Document()
    document.add_heading("Overview", level=1)
    document.add_paragraph("File Assistant paragraph")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "name"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "1"
    document.save(inputs / "report.docx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["name", "value"])
    sheet.append(["alpha", 1])
    workbook.save(inputs / "report.xlsx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "File Assistant slide"
    presentation.save(inputs / "report.pptx")

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with (inputs / "report.pdf").open("wb") as stream:
        writer.write(stream)

    media_types = {
        "report.docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "report.xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "report.pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "report.pdf": "application/pdf",
    }
    items = tuple(
        RunFileInput(
            uuid4(), name, (inputs / name).stat().st_size, "binary", media_type
        )
        for name, media_type in media_types.items()
    )
    return RunFileScope(uuid4(), inputs, tmp_path / "scratch", tmp_path / "outputs", items)


def _tools(scope: RunFileScope):
    return {tool.name: tool for tool in create_file_read_tools(lambda: scope)}


async def test_docx_tools_inspect_read_and_extract_table(tmp_path: Path) -> None:
    tools = _tools(_scope(tmp_path))
    inspected = json.loads(await tools["inspect_docx"].ainvoke({"file": "report.docx"}))
    read = json.loads(await tools["read_docx_paragraphs"].ainvoke({
        "file": "report.docx", "start": 0, "limit": 10, "max_chars": 1000,
    }))
    table = json.loads(await tools["extract_docx_tables"].ainvoke({
        "file": "report.docx", "table_index": 0, "max_rows": 10,
    }))

    assert inspected["paragraphs"] == 2
    assert read["paragraphs"][1]["text"] == "File Assistant paragraph"
    assert table["rows"][1] == ["alpha", "1"]


async def test_workbook_tool_reads_only_requested_range(tmp_path: Path) -> None:
    tools = _tools(_scope(tmp_path))
    inspected = json.loads(await tools["inspect_workbook"].ainvoke({"file": "report.xlsx"}))
    result = json.loads(await tools["read_sheet_range"].ainvoke({
        "file": "report.xlsx", "sheet": "Data", "start_row": 2, "end_row": 2,
        "start_column": 1, "end_column": 2,
    }))

    assert inspected["sheets"] == ["Data"]
    assert result["rows"] == [["alpha", 1.0]]


async def test_presentation_and_pdf_tools_use_one_based_locations(tmp_path: Path) -> None:
    tools = _tools(_scope(tmp_path))
    slides = json.loads(await tools["inspect_presentation"].ainvoke({"file": "report.pptx"}))
    slide = json.loads(await tools["read_slide"].ainvoke({
        "file": "report.pptx", "slide": 1, "max_chars": 1000,
    }))
    pdf = json.loads(await tools["inspect_pdf"].ainvoke({"file": "report.pdf"}))
    page = json.loads(await tools["read_pdf_pages"].ainvoke({
        "file": "report.pdf", "start_page": 1, "end_page": 1, "max_chars": 1000,
    }))

    assert slides["slides"] == 1
    assert slide["texts"] == ["File Assistant slide"]
    assert pdf["pages"] == 1
    assert page["pages"] == [{"page": 1, "text": ""}]


async def test_file_read_tools_publish_budget_and_trace_metadata(tmp_path: Path) -> None:
    tools = _tools(_scope(tmp_path))
    tool = tools["read_pdf_pages"]
    payload = json.loads(await tool.ainvoke({
        "file": "report.pdf", "start_page": 1, "end_page": 1,
    }))

    assert tool.metadata["side_effect_class"] == "read_only"
    assert tool.metadata["file_scope_required"] is True
    assert tool.metadata["usage_json_fields"] == {
        "scanned_bytes": "bytes_scanned",
        "returned_bytes": "bytes_returned_to_model",
    }
    assert payload["returned_bytes"] > 0
    assert payload["scanned_bytes"] > 0


async def test_wrong_format_and_missing_scope_fail_closed(tmp_path: Path) -> None:
    tools = _tools(_scope(tmp_path))
    try:
        await tools["inspect_pdf"].ainvoke({"file": "report.docx"})
    except LookupError as exc:
        assert "no pdf adapter" in str(exc)
    else:
        raise AssertionError("format-specific tool must reject a different file type")

    unavailable = {
        tool.name: tool for tool in create_file_read_tools(lambda: None)
    }
    try:
        await unavailable["inspect_pdf"].ainvoke({"file": "report.pdf"})
    except ValueError as exc:
        assert "scope is unavailable" in str(exc)
    else:
        raise AssertionError("file tool must fail closed without an active Run scope")


async def test_tool_result_fails_bounded_when_adapter_payload_is_too_large(
    tmp_path: Path,
) -> None:
    scope = _scope(tmp_path)
    registry = FileAdapterRegistry()
    registry.register(
        "pdf", SimpleNamespace(inspect=lambda resource: {"content": "x" * 600_000}),
        extensions=("pdf",),
    )
    tools = {
        tool.name: tool
        for tool in create_file_read_tools(lambda: scope, registry)
    }
    result = json.loads(await tools["inspect_pdf"].ainvoke({"file": "report.pdf"}))

    assert result["truncated"] is True
    assert "smaller range" in result["error"]
    assert result["returned_bytes"] < 512 * 1024


async def test_fast_text_deterministically_routes_complex_document(tmp_path: Path) -> None:
    scope = _scope(tmp_path)

    class Router:
        calls = []

        def should_normalize(self, resource):
            return resource.logical_name.endswith(".docx")

        async def normalize(self, account_id, run_id, resource):
            self.calls.append((account_id, run_id, resource.file_id))
            return {
                "schema_version": 1, "run_id": run_id,
                "file_id": str(resource.file_id),
                "document_ref": f"normalized-document:{run_id}:{resource.file_id}",
                "block_count": 2, "table_count": 1, "truncated": False,
            }

    router = Router()
    tools = {tool.name: tool for tool in create_file_read_tools(
        lambda: scope, document_router=router,
        account_id_provider=lambda: "account-1",
    )}
    result = json.loads(await tools["fast_text_view"].ainvoke({
        "file": "report.docx"
    }))
    assert result["route"] == "normalized_document"
    assert set(result["normalized_document_ref"]) == {
        "schema_version", "run_id", "file_id", "document_ref",
        "block_count", "table_count", "truncated",
    }
    assert router.calls == [("account-1", str(scope.run_id), scope.inputs[0].file_id)]


async def test_small_text_keeps_bounded_direct_read(tmp_path: Path) -> None:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    (inputs / "small.txt").write_text("small text", encoding="utf-8")
    item = RunFileInput(uuid4(), "small.txt", 10, "utf-8", "text/plain")
    scope = RunFileScope(
        uuid4(), inputs, tmp_path / "scratch", tmp_path / "outputs", (item,)
    )

    class Router:
        def should_normalize(self, resource):
            return resource.size_bytes > 1024

        async def normalize(self, *args):
            raise AssertionError("small text must not reach Document Worker")

    tools = {tool.name: tool for tool in create_file_read_tools(
        lambda: scope, document_router=Router(), account_id_provider=lambda: "account-1"
    )}
    result = json.loads(await tools["fast_text_view"].ainvoke({"file": "small.txt"}))
    assert result["route"] == "bounded_direct_read"
    assert "small text" in result["text"]
